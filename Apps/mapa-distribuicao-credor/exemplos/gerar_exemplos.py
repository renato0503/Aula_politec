"""Gera os dados sintéticos do Mapa de Distribuição por Credor (Sprint 0).

Saídas (nesta pasta):
  distribuicao-orgaos.xlsx        planilha de distribuição por órgão (anexo do TR)
  relatorio-pregoeiro.xlsx        resultado por item: situação, empresa, CNPJ, valor homologado
  dotacao.xlsx                    dotação por órgão, tipo (3 ou 4) e ano
  mapa-manual-referencia.xlsx     gabarito calculado aqui, de forma independente do app
  *-com-divergencias.xlsx         as mesmas planilhas com as divergências D01 a D08 plantadas
  DIVERGENCIAS.md                 lista das divergências plantadas
  seed.json                       os mesmos dados no formato do app (embutido no index.html)

Regras do cálculo (PROJETO.md, RN01 a RN11):
  total do item = quantidade × valor unitário, arredondado a 2 casas (RN01, RN09)
  homologado substitui o estimado; fracassado e deserto ficam sem valor (RN03)
  dotação consumida na ordem dos grupos; o excedente vai para o ano seguinte (RN10)

Tudo é fictício. Semente fixa = 20261029.
Uso: python "Apps/mapa-distribuicao-credor/exemplos/gerar_exemplos.py"
"""
import json
import pathlib
import random
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

AQUI = pathlib.Path(__file__).resolve().parent
random.seed(20261029)

PROCESSO = {"id": "PROC-2026-001", "nome": "Aquisição de insumos e equipamentos de laboratório", "numero": "001/2026", "ano": 2026}
ORGAOS = ["POLITEC-SEDE", "IML-CUIABA", "LAB-FORENSE"]
ANO = 2026

GRUPOS = [  # grupo, tipo de despesa, situação
    ("G01", 3, "homologado"),
    ("G02", 3, "homologado"),
    ("G03", 4, "fracassado"),
    ("G04", 4, "deserto"),
    ("G05", 3, "homologado"),  # misto: itens com empresas diferentes
    ("G06", 4, "homologado"),  # misto
]
CATALOGO = {
    3: [("Luva nitrílica tamanho M", "caixa"), ("Reagente luminol 500 mL", "frasco"), ("Swab estéril", "pacote"),
        ("Etiqueta de lacre numerada", "rolo"), ("Envelope de segurança", "pacote"), ("Álcool isopropílico 1 L", "frasco"),
        ("Máscara PFF2", "caixa"), ("Tubo de coleta 10 mL", "pacote"), ("Papel filtro qualitativo", "caixa"),
        ("Ponteira 1000 µL", "caixa"), ("Fita adesiva de evidência", "rolo"), ("Saco plástico de evidência", "pacote"),
        ("Óculos de proteção", "unidade"), ("Avental descartável", "pacote"), ("Reagente de Kastle-Meyer", "frasco")],
    4: [("Microscópio estereoscópico", "unidade"), ("Centrífuga de bancada", "unidade"), ("Balança analítica", "unidade"),
        ("Câmera fotográfica forense", "unidade"), ("Luz forense multibanda", "unidade"), ("Capela de fluxo laminar", "unidade"),
        ("Freezer -20 °C", "unidade"), ("Computador de extração forense", "unidade"), ("Bloqueador de escrita", "unidade"),
        ("Pipeta automática", "unidade"), ("Estufa de secagem", "unidade"), ("Agitador vórtex", "unidade"),
        ("Leitor de microplacas", "unidade"), ("Destilador de água", "unidade"), ("Termociclador", "unidade")],
}


def cnpj_valido(base12):
    def dv(nums, pesos):
        s = sum(n * p for n, p in zip(nums, pesos)) % 11
        return 0 if s < 2 else 11 - s
    n = [int(c) for c in base12]
    d1 = dv(n, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    d2 = dv(n + [d1], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    return base12 + str(d1) + str(d2)


EMPRESAS = [
    {"empresa": "Alfa Suprimentos Laboratoriais Ltda", "cnpj": cnpj_valido("112223330001")},
    {"empresa": "Beta Científica Comércio Ltda", "cnpj": cnpj_valido("445556660001")},
    {"empresa": "Gama Equipamentos Técnicos S.A.", "cnpj": cnpj_valido("778889990001")},
    {"empresa": "Delta Distribuidora Hospitalar Ltda", "cnpj": cnpj_valido("223334440001")},
]


def dinheiro(v):
    return Decimal(v).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def gerar():
    itens, distribuicao, resultado = [], [], []
    usados = {3: 0, 4: 0}
    for g_idx, (grupo, tipo, situacao) in enumerate(GRUPOS):
        for i in range(5):
            desc, unidade = CATALOGO[tipo][usados[tipo]]
            usados[tipo] += 1
            item = f"{i + 1:02d}"
            vu = dinheiro(random.uniform(8, 180) if tipo == 3 else random.uniform(2500, 48000))
            qtds = {o: (random.randint(5, 80) if tipo == 3 else random.randint(1, 4)) for o in ORGAOS}
            itens.append({"grupo": grupo, "item": item, "codigo": f"CATMAT-{4100 + g_idx * 10 + i}", "descricao": desc,
                          "unidade": unidade, "tipo_despesa": tipo, "valor_unitario_estimado": vu,
                          "quantidade_total": sum(qtds.values())})
            for o, q in qtds.items():
                distribuicao.append({"orgao": o, "grupo": grupo, "item": item, "quantidade": q})
            if situacao == "homologado":
                if grupo in ("G05", "G06"):
                    emp = EMPRESAS[(i % 2) + (2 if grupo == "G06" else 0)]
                else:
                    emp = EMPRESAS[0 if grupo == "G01" else 1]
                desconto = Decimal(str(round(random.uniform(0.04, 0.22), 4)))
                vh = dinheiro(vu * (1 - desconto))
                resultado.append({"grupo": grupo, "item": item, "situacao": situacao, "empresa": emp["empresa"],
                                  "cnpj": emp["cnpj"], "valor_unitario_homologado": vh})
            else:
                resultado.append({"grupo": grupo, "item": item, "situacao": situacao, "empresa": "", "cnpj": "",
                                  "valor_unitario_homologado": None})
    return itens, distribuicao, resultado


def calcular(itens, distribuicao, resultado, dotacao):
    """Gabarito: implementação independente do app, em Decimal."""
    res = {(r["grupo"], r["item"]): r for r in resultado}
    tipo_grupo = {g: t for g, t, _ in GRUPOS}
    mapa_item, por_orgao = [], []
    for it in itens:
        r = res[(it["grupo"], it["item"])]
        total_est = dinheiro(it["quantidade_total"] * it["valor_unitario_estimado"])
        vh = r["valor_unitario_homologado"]
        total_hom = dinheiro(it["quantidade_total"] * vh) if vh is not None else None
        mapa_item.append({**it, "total_estimado": total_est, "situacao": r["situacao"], "empresa": r["empresa"],
                          "valor_unitario_homologado": vh, "total_homologado": total_hom})
    for d in distribuicao:
        it = next(x for x in itens if x["grupo"] == d["grupo"] and x["item"] == d["item"])
        r = res[(d["grupo"], d["item"])]
        vh = r["valor_unitario_homologado"]
        por_orgao.append({**d, "valor_estimado": dinheiro(d["quantidade"] * it["valor_unitario_estimado"]),
                          "valor_homologado": dinheiro(d["quantidade"] * vh) if vh is not None else None,
                          "empresa": r["empresa"], "cnpj": r["cnpj"], "tipo_despesa": tipo_grupo[d["grupo"]]})
    # empresas x órgão
    emp_orgao = {}
    for p in por_orgao:
        if p["valor_homologado"] is None:
            continue
        k = (p["empresa"], p["cnpj"], p["orgao"])
        emp_orgao[k] = emp_orgao.get(k, Decimal("0")) + p["valor_homologado"]
    # distribuição orçamentária: ordem dos grupos, depois dos itens (RN10)
    saldo = {(d["orgao"], d["tipo"]): d["valor"] for d in dotacao}
    orc = {}
    for p in sorted(por_orgao, key=lambda x: (x["grupo"], x["item"], x["orgao"])):
        if p["valor_homologado"] is None:
            continue
        k_saldo = (p["orgao"], p["tipo_despesa"])
        disp = saldo.get(k_saldo, Decimal("0"))
        no_ano = min(disp, p["valor_homologado"])
        saldo[k_saldo] = disp - no_ano
        k = (p["cnpj"], p["empresa"], p["orgao"], p["tipo_despesa"])
        a = orc.setdefault(k, [Decimal("0"), Decimal("0")])
        a[0] += no_ano
        a[1] += p["valor_homologado"] - no_ano
    return mapa_item, por_orgao, emp_orgao, orc


def gerar_dotacao(por_orgao_hom):
    """Dotação suficiente, exceto investimento do LAB-FORENSE (força o ano seguinte, RN10)."""
    out = []
    for o in ORGAOS:
        for t in (3, 4):
            precisa = sum((p["valor_homologado"] for p in por_orgao_hom
                           if p["orgao"] == o and p["tipo_despesa"] == t and p["valor_homologado"] is not None), Decimal("0"))
            valor = dinheiro(precisa * Decimal("0.6")) if (o == "LAB-FORENSE" and t == 4) else dinheiro(precisa * Decimal("1.15") + 1000)
            out.append({"orgao": o, "tipo": t, "ano": ANO, "valor": valor})
    return out


CAB = Font(bold=True)
FUNDO = PatternFill("solid", fgColor="DDE7F5")


def planilha(caminho, abas):
    wb = Workbook()
    wb.remove(wb.active)
    for nome, cab, linhas in abas:
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill = CAB, FUNDO
        for l in linhas:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in l])
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(caminho)


def salvar_entradas(sufixo, itens, distribuicao, resultado, dotacao):
    q = {(d["orgao"], d["grupo"], d["item"]): d["quantidade"] for d in distribuicao}
    planilha(AQUI / f"distribuicao-orgaos{sufixo}.xlsx", [(
        "Distribuicao",
        ["Grupo", "Item", "Código", "Descrição", "Unidade", "Tipo de despesa", "Valor unitário estimado", "Quantidade total"] + ORGAOS,
        [[i["grupo"], i["item"], i["codigo"], i["descricao"], i["unidade"], i["tipo_despesa"], i["valor_unitario_estimado"],
          i["quantidade_total"]] + [q.get((o, i["grupo"], i["item"]), 0) for o in ORGAOS] for i in itens])])
    planilha(AQUI / f"relatorio-pregoeiro{sufixo}.xlsx", [(
        "Resultado", ["Grupo", "Item", "Situação", "Empresa", "CNPJ", "Valor unitário homologado"],
        [[r["grupo"], r["item"], r["situacao"], r["empresa"], r["cnpj"], r["valor_unitario_homologado"]] for r in resultado])])
    planilha(AQUI / f"dotacao{sufixo}.xlsx", [(
        "Dotacao", ["Órgão", "Tipo", "Ano", "Valor"], [[d["orgao"], d["tipo"], d["ano"], d["valor"]] for d in dotacao])])


def main():
    itens, distribuicao, resultado = gerar()
    _, por_orgao, _, _ = calcular(itens, distribuicao, resultado, [])
    dotacao = gerar_dotacao(por_orgao)
    mapa_item, por_orgao, emp_orgao, orc = calcular(itens, distribuicao, resultado, dotacao)

    salvar_entradas("", itens, distribuicao, resultado, dotacao)
    planilha(AQUI / "mapa-manual-referencia.xlsx", [
        ("1 Mapa por item", ["Grupo", "Item", "Descrição", "Quantidade", "Valor unitário estimado", "Total estimado",
                             "Situação", "Empresa", "Valor unitário homologado", "Total homologado"],
         [[m["grupo"], m["item"], m["descricao"], m["quantidade_total"], m["valor_unitario_estimado"], m["total_estimado"],
           m["situacao"], m["empresa"], m["valor_unitario_homologado"], m["total_homologado"]] for m in mapa_item]),
        ("2 Por órgão", ["Órgão", "Grupo", "Item", "Quantidade", "Valor estimado", "Valor homologado"],
         [[p["orgao"], p["grupo"], p["item"], p["quantidade"], p["valor_estimado"], p["valor_homologado"]] for p in por_orgao]),
        ("6 Empresa x órgão", ["Empresa", "CNPJ", "Órgão", "Valor homologado"],
         [[e, c, o, v] for (e, c, o), v in sorted(emp_orgao.items())]),
        ("8 Orçamentária", ["CNPJ", "Empresa", "Órgão", "Tipo", f"Valor {ANO}", f"Valor {ANO + 1}"],
         [[c, e, o, t, v[0], v[1]] for (c, e, o, t), v in sorted(orc.items())]),
        ("Totais", ["Indicador", "Valor"],
         [["Total estimado", sum(m["total_estimado"] for m in mapa_item)],
          ["Total homologado", sum(m["total_homologado"] for m in mapa_item if m["total_homologado"] is not None)],
          [f"Orçamentário {ANO}", sum(v[0] for v in orc.values())],
          [f"Orçamentário {ANO + 1}", sum(v[1] for v in orc.values())]]),
    ])

    # variante com divergências plantadas
    it_d = json.loads(json.dumps(itens, default=str))
    di_d = json.loads(json.dumps(distribuicao, default=str))
    re_d = json.loads(json.dumps(resultado, default=str))
    do_d = json.loads(json.dumps(dotacao, default=str))
    plantadas = []
    it_d[1]["quantidade_total"] += 3
    plantadas.append(("D01", f"G01 item 02: quantidade total {it_d[1]['quantidade_total']} difere da soma por órgão"))
    di_d.append({"orgao": "IML-CUIABA", "grupo": "G07", "item": "01", "quantidade": 2})
    plantadas.append(("D02", "Distribuição traz G07 item 01, que não existe no relatório do pregoeiro"))
    re_d[2]["valor_unitario_homologado"] = str(dinheiro(Decimal(it_d[2]["valor_unitario_estimado"]) * Decimal("1.08")))
    plantadas.append(("D03", "G01 item 03: homologado 8% acima do estimado"))
    re_d[5]["cnpj"] = ""
    plantadas.append(("D04", "G02 item 01: homologado sem CNPJ"))
    re_d[21]["cnpj"] = re_d[21]["cnpj"][:-1] + str((int(re_d[21]["cnpj"][-1]) + 1) % 10)
    plantadas.append(("D05", f"G05 item 02: CNPJ {re_d[21]['cnpj']} com dígito verificador inválido"))
    plantadas.append(("D06", "Consequência de D01: a soma por órgão de G01 não iguala o total do grupo"))
    do_d = [d for d in do_d if not (d["orgao"] == "IML-CUIABA" and d["tipo"] == 4)]
    plantadas.append(("D07", "IML-CUIABA sem dotação do tipo 4 (investimento)"))
    plantadas.append(("D08", "LAB-FORENSE tipo 4: dotação cobre 60% e o restante vai para o ano seguinte (também no gabarito)"))
    salvar_entradas("-com-divergencias", it_d, di_d, re_d, do_d)
    (AQUI / "DIVERGENCIAS.md").write_text(
        "# Divergências plantadas\n\nNos arquivos `*-com-divergencias.xlsx`. Os arquivos sem sufixo são o caso limpo, "
        "que precisa bater com `mapa-manual-referencia.xlsx` com diferença zero.\n\n| Código | Onde |\n|---|---|\n" +
        "".join(f"| {c} | {d} |\n" for c, d in plantadas), encoding="utf-8")

    def c(v):  # dinheiro em centavos inteiros (formato do app)
        return None if v is None else int((Decimal(v) * 100).to_integral_value())
    seed = {
        "processos": [{**PROCESSO, "status": "em andamento", "criado_em": "2026-10-29T12:00:00Z"}],
        "itens": [{"id": f"{PROCESSO['id']}|{i['grupo']}|{i['item']}", "id_processo": PROCESSO["id"], "grupo": i["grupo"], "item": i["item"],
                   "codigo": i["codigo"], "descricao": i["descricao"], "unidade": i["unidade"], "tipo_despesa": i["tipo_despesa"],
                   "valor_unitario_estimado_c": c(i["valor_unitario_estimado"]), "quantidade_total": i["quantidade_total"],
                   "origem": {"arquivo": "distribuicao-orgaos.xlsx", "aba": "Distribuicao", "linha": n + 2}} for n, i in enumerate(itens)],
        "distribuicao": [{"id": f"{PROCESSO['id']}|{d['orgao']}|{d['grupo']}|{d['item']}", "id_processo": PROCESSO["id"], **d} for d in distribuicao],
        "resultado": [{"id": f"{PROCESSO['id']}|{r['grupo']}|{r['item']}", "id_processo": PROCESSO["id"], "grupo": r["grupo"], "item": r["item"],
                       "situacao": r["situacao"], "empresa": r["empresa"], "cnpj": r["cnpj"], "valor_unitario_homologado_c": c(r["valor_unitario_homologado"]),
                       "origem": {"arquivo": "relatorio-pregoeiro.xlsx", "aba": "Resultado", "linha": n + 2}} for n, r in enumerate(resultado)],
        "dotacao": [{"id": f"{PROCESSO['id']}|{d['orgao']}|{d['tipo']}|{d['ano']}", "id_processo": PROCESSO["id"], "orgao": d["orgao"], "tipo": d["tipo"],
                     "ano": d["ano"], "valor_c": c(d["valor"])} for d in dotacao],
        "gabarito": [{"id": "totais", "total_estimado_c": c(sum(m["total_estimado"] for m in mapa_item)),
                      "total_homologado_c": c(sum(m["total_homologado"] for m in mapa_item if m["total_homologado"] is not None)),
                      "orcamento_ano_c": c(sum(v[0] for v in orc.values())), "orcamento_ano_seguinte_c": c(sum(v[1] for v in orc.values()))}],
    }
    (AQUI / "seed.json").write_text(json.dumps(seed, ensure_ascii=False, indent=1), encoding="utf-8")
    print("itens", len(itens), "| distribuição", len(distribuicao), "| resultado", len(resultado), "| dotação", len(dotacao))
    print("gabarito", seed["gabarito"][0])


if __name__ == "__main__":
    main()
