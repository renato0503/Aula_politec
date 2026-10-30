"""Cenários de interface de cada sprint (o "Pronto quando" do sprint.md de cada app).

Cada cenário abre o index.html no Chromium, usa a tela como uma pessoa usaria
(clicar, preencher, escolher arquivo) e confere o resultado no banco do app.
Complementa o testar.py, que confere a estrutura (rotas, testes embutidos, rede).

Uso: python Apps/_base/testar_sprints.py [app ...]
"""
import pathlib
import sys
import tempfile

from playwright.sync_api import sync_playwright

APPS = pathlib.Path(__file__).resolve().parent.parent
TMP = pathlib.Path(tempfile.mkdtemp(prefix="sprints-"))
CENARIOS = {}


def cenario(app):
    def reg(fn):
        CENARIOS.setdefault(app, []).append(fn)
        return fn
    return reg


def contar(pg, store):
    return pg.evaluate("async s => (await db.list(s)).length", store)


def ir(pg, url, rota, espera=250):
    pg.goto(url + rota)
    pg.wait_for_timeout(espera)


def perfil(pg, url, nome):
    ir(pg, url, "#/config")
    pg.evaluate("async p => { await db.meta('perfil', p); }", nome)
    pg.reload()
    # espera o IndexedDB abrir: com versao nova do banco a abertura pode passar de 300 ms
    pg.wait_for_function("() => typeof db !== 'undefined' && !!db.con", timeout=15000)
    pg.wait_for_timeout(300)


# ------------------------------------------------------------------ mapa (Igor)
@cenario("mapa-distribuicao-credor")
def s2_importacao(pg, url, ok):
    ex = APPS / "mapa-distribuicao-credor" / "exemplos"
    ir(pg, url, "#/processos")
    pg.fill("#procNumero", "001/2026"); pg.fill("#procNome", "Teste de importação"); pg.click("#formProcesso button[type=submit]")
    pg.wait_for_timeout(400)
    ok("criar processo leva à importação", "#/importar/" in pg.url)
    for fonte, arq in [("distribuicao", "distribuicao-orgaos.xlsx"), ("pregoeiro", "relatorio-pregoeiro.xlsx"), ("dotacao", "dotacao.xlsx")]:
        pg.set_input_files(f"input[data-fonte={fonte}]", str(ex / arq)); pg.wait_for_timeout(300)
    pg.click("#btnImportar"); pg.wait_for_timeout(1200)
    ok("30 itens", contar(pg, "itens") == 30)
    ok("90 linhas por órgão", contar(pg, "distribuicao") == 90)
    ok("30 resultados", contar(pg, "resultado") == 30)
    ok("6 dotações", contar(pg, "dotacao") == 6)
    ok("origem com arquivo e linha", pg.evaluate("async ()=> (await db.list('itens'))[0].origem.linha > 1"))
    ok("totais estimados batem com o gabarito do exemplo", pg.evaluate("""async ()=>{
        const i = await db.list('itens'); const g = dadosEmbutidos('seed').gabarito[0];
        return i.reduce((a, x)=>a + dinheiro.vezes(x.quantidade_total, x.valor_unitario_estimado_c), 0) === g.total_estimado_c; }"""))
    # sem a coluna de quantidade: recusa com a mensagem certa
    from openpyxl import load_workbook
    wb = load_workbook(ex / "distribuicao-orgaos.xlsx"); ws = wb.active
    ws.delete_cols(8)  # "Quantidade total"
    sem = TMP / "sem-quantidade.xlsx"; wb.save(sem)
    pg.set_input_files("input[data-fonte=distribuicao]", str(sem)); pg.wait_for_timeout(300)
    pg.select_option("#map-distribuicao-quantidade_total", "-1")
    pg.click("#btnImportar"); pg.wait_for_timeout(800)
    ok("planilha sem quantidade é recusada com a mensagem certa", "Quantidade total" in pg.inner_text("#erros-distribuicao"))
    ok("a recusa não apagou os itens importados antes", contar(pg, "itens") == 30)
    # com divergências: CNPJ inválido e sem CNPJ viram alertas
    pg.set_input_files("input[data-fonte=pregoeiro]", str(ex / "relatorio-pregoeiro-com-divergencias.xlsx")); pg.wait_for_timeout(300)
    pg.click("#btnImportar"); pg.wait_for_timeout(800)
    rel = pg.inner_text("#rel-pregoeiro")
    ok("CNPJ inválido e homologado sem CNPJ aparecem como alerta", "inválido" in rel and "sem CNPJ" in rel)


# ------------------------------------------------------------------ equipamentos (João)
@cenario("equipamentos-google-earth")
def s2_unidades(pg, url, ok):
    perfil(pg, url, "coordenadoria")
    ir(pg, url, "#/config"); pg.click("#btnSeed"); pg.wait_for_timeout(800)
    ir(pg, url, "#/unidades", 500)
    ok("as 12 unidades do exemplo aparecem com coordenadas", pg.locator("table tr").count() == 13 and "-15." in pg.inner_text("table"))
    dir_id = pg.evaluate("async ()=> (await db.list('diretorias'))[0].id")
    pg.fill("#uniNome", "Posto Teste"); pg.select_option("#uniDiretoria", dir_id)
    pg.fill("#uniMunicipio", "Cuiabá"); pg.dispatch_event("#uniMunicipio", "change")
    ok("o município preenche as coordenadas", pg.input_value("#uniLat").startswith("-15"))
    pg.fill("#uniLat", "-16.68"); pg.fill("#uniLon", "-49.25")  # Goiânia
    pg.click("#formUnidade button[type=submit]"); pg.wait_for_timeout(400)
    ok("coordenada de Goiânia é recusada (RN08)", "fora de Mato Grosso" in pg.inner_text("#errosUnidade"))
    ok("nada foi gravado", contar(pg, "unidades") == 12)
    pg.fill("#uniLat", "-15.60"); pg.fill("#uniLon", "-56.10")
    pg.click("#formUnidade button[type=submit]"); pg.wait_for_timeout(600)
    ok("coordenada de Cuiabá é aceita", contar(pg, "unidades") == 13)
    pg.locator("[data-excluir]").first.click(); pg.wait_for_timeout(500)
    ok("excluir é lógico (o registro continua, marcado)", contar(pg, "unidades") == 13 and pg.evaluate("async ()=> (await db.list('unidades')).filter(u=>u.excluido).length") == 1)
    ok("a unidade nova tem id, datas e município IBGE", pg.evaluate("async ()=>{ const u = (await db.list('unidades')).find(x=>x.nome==='Posto Teste'); return !!(u && u.criadoEm && u.municipio_ibge === '5103403'); }"))


# ------------------------------------------------------------------ inscrições (Edilaura)
@cenario("inscricoes-cursos")
def s2_curso_e_link(pg, url, ok):
    perfil(pg, url, "gerencia")
    ir(pg, url, "#/config"); pg.click("#btnSeed"); pg.wait_for_timeout(600)
    ir(pg, url, "#/cursos")
    pg.fill("#cNome", "Perícia em Local de Crime"); pg.fill("#cCarga", "16"); pg.fill("#cLocal", "Auditório")
    pg.fill("#cInicio", "2026-12-07"); pg.fill("#cFim", "2026-12-08"); pg.fill("#cInscIni", "2026-11-01"); pg.fill("#cInscFim", "2026-11-20")
    pg.fill("#cTriagem", "2026-11-25"); pg.fill("#cGrupo", "https://chat.exemplo/turma")
    pg.click("#formCurso button[type=submit]"); pg.wait_for_timeout(400)
    ok("curso sem vagas é recusado", "ao menos uma vaga" in pg.inner_text("#errosCurso"))
    pg.fill("#vaga-DMC", "5"); pg.fill("#vaga-DIN", "3"); pg.dispatch_event("#vaga-DIN", "input")
    ok("total de vagas soma as diretorias", pg.inner_text("#totalVagas") == "8")
    pg.click("#formCurso button[type=submit]"); pg.wait_for_timeout(600)
    ok("salvar leva à divulgação", "#/divulgar/" in pg.url)
    link = pg.inner_text("#linkInscricao")
    ok("QR Code gerado", pg.locator("#qr svg").count() == 1)
    ok("texto de divulgação com as vagas", "5 vagas" in pg.input_value("#textoDivulgacao"))
    novo = pg.context.browser.new_context(); p2 = novo.new_page()   # outro navegador, sem dados
    p2.goto(link); p2.wait_for_timeout(500)
    ok("o link abre a inscrição com o curso, sem dado prévio", "Perícia em Local de Crime" in p2.inner_text("#conteudo"))
    ok("a inscrição não mostra o menu da gerência", p2.locator("#menu").is_hidden())
    p2.goto(link[:-1] + ("1" if link.endswith("0") else "0")); p2.wait_for_timeout(300)
    ok("link alterado mostra aviso", "alterado" in p2.inner_text("#conteudo"))
    novo.close()


# ------------------------------------------------------------------ Hórus (Iohan)
@cenario("horus-banco-demandas")
def s2_cadastro_demanda(pg, url, ok):
    perfil(pg, url, "USUARIO_SETORIAL")
    ir(pg, url, "#/config"); pg.click("#btnSeed"); pg.wait_for_timeout(700); ir(pg, url, "#/config", 400)
    pg.fill("#cfgUsuario", "teste.usuario"); pg.dispatch_event("#cfgUsuario", "change")
    uni = pg.evaluate("async ()=> (await db.list('unidades'))[0].id"); pg.select_option("#selUnidade", uni); pg.wait_for_timeout(200)
    ir(pg, url, "#/demanda/nova", 400)
    pg.fill("#dValor", "33,33"); pg.fill("#dQtd", "3")
    medida = pg.evaluate("async ()=> (await db.list('unidades_medida'))[0].id"); pg.select_option("#dMedida", medida)
    pg.fill("#dEspec", "Reagente de teste para o laboratório, frasco de 500 mL, validade de 12 meses.")
    pg.fill("#dJust", "x" * 99); pg.fill("#dLinks", "https://exemplo.com/produto")
    ok("valor global calculado na hora (R4)", pg.inner_text("#dGlobal") == "R$ 99,99")
    n0 = contar(pg, "banco_demandas")
    pg.click("#formDemanda button[type=submit]"); pg.wait_for_timeout(400)
    ok("justificativa de 99 caracteres é recusada (R2)", "R2" in pg.inner_text("#errosDemanda") and contar(pg, "banco_demandas") == n0)
    pg.fill("#dJust", "x" * 100)
    pg.set_input_files("#dPdf", {"name": "orcamento.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4 teste"})
    pg.click("#formDemanda button[type=submit]"); pg.wait_for_timeout(500)
    ok("com 100 caracteres, a demanda é salva", contar(pg, "banco_demandas") == n0 + 1)
    d = pg.evaluate("async ()=> (await db.list('banco_demandas')).find(x=>x.criado_por==='teste.usuario')")
    ok("valor global gravado = 99,99", d and d["valor_global"] == 99.99)
    ok("PDF guardado com a demanda", pg.evaluate("async ()=> (await db.list('banco_demandas')).find(x=>x.criado_por==='teste.usuario').arquivo_pdf instanceof Blob"))
    alheia = pg.evaluate("async ()=> (await db.list('banco_demandas')).find(x=>x.criado_por!=='teste.usuario').id")
    ir(pg, url, "#/demanda/" + alheia, 400)
    ok("demanda de outro autor não abre para edição (R8)", "R8" in pg.inner_text("#conteudo") and pg.locator("#formDemanda").count() == 0)
    ir(pg, url, "#/demanda/" + d["id"], 400)
    ok("o autor edita a própria demanda", pg.locator("#formDemanda").count() == 1)


# ------------------------------------------------------------------ Agiliza + (grupo)
@cenario("agiliza-mais")
def s2_pasta_do_caso(pg, url, ok):
    ex = APPS / "agiliza-mais" / "exemplos"
    def inventario(caso):
        ir(pg, url, "#/inicio", 300); pg.click("#btnNovoCaso"); pg.wait_for_timeout(400)
        pg.set_input_files("#inpPasta", str(ex / caso)); pg.wait_for_timeout(800)
        return pg.evaluate("async ()=> (await db.get('casos', await db.meta('caso_atual'))).inventario")
    inv = inventario("caso-ficticio-001")
    st = {i["chave"]: i["status"] for i in inv["itens"]}
    ok("caso 001: requisição, device e hash presentes", st.get("requisicao") == st.get("device") == st.get("hash") == "presente")
    ok("caso 001: 2 cartões SIM (01-A e 01-B)", inv["qtdSims"] == 2 and "sim:01-A" in st and "sim:01-B" in st)
    ok("caso 001: cronômetro iniciado", pg.evaluate("async ()=> !!(await db.get('casos', await db.meta('caso_atual'))).inicio_em"))
    inv = inventario("caso-ficticio-002")
    ok("caso 002: sem SIM e sem erro (RN06)", inv["qtdSims"] == 0 and not inv["faltando"])
    inv = inventario("caso-ficticio-003")
    ok("caso 003: acusa o HASHDOHASH.txt ausente", inv["faltando"] == ["hash"])
    pg.click("#btnAvancar"); pg.wait_for_timeout(300)
    ok("caso 003: avançar é bloqueado (RN02)", "HASHDOHASH" in pg.inner_text("#avisoAvancar") and "#/arquivos" in pg.url)
    pg.check("[data-manual=hash]"); pg.wait_for_timeout(200); pg.click("#btnAvancar"); pg.wait_for_timeout(300)
    ok("informar manualmente libera o avanço", "#/dados" in pg.url)
    ok("3 casos criados", contar(pg, "casos") == 3)
    ok("os originais não mudaram de tamanho", (ex / "caso-ficticio-001" / "DEVICE INFO.txt").stat().st_size > 0)


# ------------------------------------------------------------------ Agendamento CIN (Elthon)
@cenario("agendamento-cin")
def s2_horarios_e_postos(pg, url, ok):
    perfil(pg, url, "posto")
    pg.evaluate("async ()=> db.meta('posto_id', 'P01')"); pg.reload(); pg.wait_for_timeout(300)
    ir(pg, url, "#/posto/horarios", 500)
    pg.fill("#gDe", "2026-11-16"); pg.fill("#gAte", "2026-11-27"); pg.click("#formGrade button[type=submit]"); pg.wait_for_timeout(1500)
    ok("gera 2 semanas de horários (160 novos)", pg.evaluate("async ()=> (await db.porIndice('horarios','posto_id','P01')).filter(h=>h.data>='2026-11-16'&&h.data<='2026-11-27').length") == 160)
    pg.click("#formGrade button[type=submit]"); pg.wait_for_timeout(1500)
    ok("gerar de novo não duplica", pg.evaluate("async ()=> (await db.porIndice('horarios','posto_id','P01')).filter(h=>h.data>='2026-11-16'&&h.data<='2026-11-27').length") == 160)
    pg.click("[data-dia-bloq='2026-11-16']"); pg.wait_for_timeout(800)
    ok("bloquear o dia bloqueia os 16 horários", pg.evaluate("async ()=> (await db.porIndice('horarios','posto_id','P01')).filter(h=>h.data==='2026-11-16'&&h.situacao==='bloqueado').length") == 16)
    pg.evaluate("async ()=> db.meta('posto_id', 'P02')"); pg.reload(); pg.wait_for_timeout(300)
    msg = pg.evaluate("async ()=>{ try { await api.gerarHorarios('P01', {de:'2026-12-01',ate:'2026-12-01',dias:[2],inicio:'08:00',fim:'09:00',intervalo:30,vagas:1}); return 'passou'; } catch(e){ return e.message; } }")
    ok("o posto P02 não consegue alterar a agenda do P01 (RN05)", "RN05" in msg)
    perfil(pg, url, "gestao")
    ir(pg, url, "#/gestao", 400)
    pg.fill("#pId", "p03"); pg.fill("#pNome", "Posto Teste"); pg.fill("#pMunicipio", "Sinop"); pg.click("#formPosto button[type=submit]"); pg.wait_for_timeout(500)
    ok("gestão cadastra posto", contar(pg, "postos") == 3)
    pg.click("[data-ativo='P03']"); pg.wait_for_timeout(500)
    ok("posto desativado sai da lista pública (RN01)", pg.evaluate("async ()=> (await api.listarPostos()).length") == 2)


# ------------------------------------------------------------------ Adiantamento (Jaqueline e Linicker)
@cenario("adiantamento")
def s2_novo_pedido(pg, url, ok):
    ir(pg, url, "#/pedido/novo", 400)
    pg.fill("#pUnidade", "IML Cuiabá"); pg.fill("#pSolicitante", "Servidor Teste"); pg.fill("#pMatricula", "123456")
    pg.fill("#pValor", "2.500,00"); pg.fill("#pFinalidade", "Conserto emergencial do gerador")
    pg.fill("#pJust", "O gerador parou e o IML ficou sem energia para as câmaras frias; o conserto é urgente.")
    pg.fill("#oForn1", "Fornecedor A"); pg.fill("#oCnpj1", "11.222.333/0001-81"); pg.fill("#oValor1", "2.400,00")
    pg.fill("#oForn2", "Fornecedor B"); pg.fill("#oCnpj2", "11222333000182"); pg.fill("#oValor2", "2.600,00")   # CNPJ inválido
    pg.click("#formPedido button[type=submit]"); pg.wait_for_timeout(400)
    erros = pg.inner_text("#errosPedido")
    ok("2 orçamentos são recusados (RN01)", "RN01" in erros)
    ok("CNPJ inválido é apontado", "Orçamento 2: CNPJ inválido" in erros)
    ok("nada foi gravado", contar(pg, "processos") == 0)
    pg.fill("#oCnpj2", "44.555.666/0001-81"); pg.fill("#oForn3", "Fornecedor C"); pg.fill("#oCnpj3", "77.888.999/0001-81"); pg.fill("#oValor3", "2.550,00")
    pg.click("#formPedido button[type=submit]"); pg.wait_for_timeout(600)
    ok("com 3 orçamentos válidos, o pedido é registrado", contar(pg, "processos") == 1 and contar(pg, "orcamentos") == 3)
    ok("o pedido entra na etapa Análise da DAS (RN02)", pg.evaluate("async ()=> (await db.list('processos'))[0].etapa_atual") == "Análise da DAS")
    ok("valor gravado sem erro de centavo", pg.evaluate("async ()=> (await db.list('processos'))[0].valor_solicitado") == 2500)
    ok("etapa da DAS aberta, sem saída", pg.evaluate("async ()=> (await db.list('etapas')).filter(e=>e.saida_em===null).length") == 1)


# ------------------------------------------------------------------ RNC (Sabrina)
@cenario("rnc-nao-conformidade")
def s2_importar_planilha(pg, url, ok):
    ex = APPS / "rnc-nao-conformidade" / "exemplos"
    perfil(pg, url, "cqual")
    pg.evaluate("async ()=>{ for (const s of dadosEmbutidos('seed').setores) await db.put('setores', s); }")
    def importar(arq):
        ir(pg, url, "#/importar", 300); pg.set_input_files("#arqRnc", str(arq)); pg.wait_for_timeout(400)
        pg.click("#btnImportarRnc"); pg.wait_for_timeout(1200)
    importar(ex / "respostas-forms-ficticio.csv")
    ok("as 20 RNCs do Forms viram 20 novas", pg.inner_text("#qtdNovas") == "20 novas" and contar(pg, "rncs") == 20)
    ok("cada RNC abre a etapa de validação com prazo de 5 dias", pg.evaluate("async ()=>{ const e = (await db.list('etapas'))[0]; return (new Date(e.prazo) - new Date(e.inicio)) / 86400000 === 5; }"))
    ok("setores casados pela sigla", pg.evaluate("async ()=> (await db.list('rncs')).every(r=>r.setor_relator.length <= 4)"))
    importar(ex / "respostas-forms-ficticio.csv")
    ok("importar de novo não duplica", pg.inner_text("#qtdExistentes") == "20 já existentes" and contar(pg, "rncs") == 20)
    # planilha com nome do servidor e sem a coluna Número
    linhas = (ex / "respostas-forms-ficticio.csv").read_text(encoding="utf-8-sig").splitlines()
    cab = linhas[0].split(",")
    nova = [",".join(cab[:-1] + ["Nome do servidor"])] + [",".join(l.split(",")[:-1] + ["Fulano de Tal"]).replace("2026,", "2026,").replace("/09/2026", "/10/2026", 1) for l in linhas[1:3]]
    arq = TMP / "forms-com-nome.csv"; arq.write_text("\n".join(nova), encoding="utf-8")
    importar(arq)
    ok("coluna Nome do servidor é descartada (RN07)", "Nome do servidor" in pg.inner_text("#mapaRnc") and "descartadas" in pg.inner_text("#mapaRnc"))
    ok("nenhum nome de servidor gravado", "Fulano" not in pg.evaluate("async ()=> JSON.stringify(await db.list('rncs'))"))
    ok("sem coluna Número, o app numera DLF/2026-021 em diante", pg.evaluate("async ()=> !!(await db.get('rncs', 'DLF/2026-021'))"))


# ------------------------------------------------------------------ Metas BI (Karine)
@cenario("metas-bi-ocr")
def s2_ler_csv_excel(pg, url, ok):
    ex = APPS / "metas-bi-ocr" / "exemplos"
    ir(pg, url, "#/config"); pg.click("#btnSeed"); pg.wait_for_timeout(600)
    for mes in ["2026-07", "2026-08", "2026-09"]:
        ir(pg, url, "#/importar", 300); pg.set_input_files("#arqMes", str(ex / f"powerbi1-{mes}.csv")); pg.wait_for_timeout(800)
    ok("3 meses importados", contar(pg, "relatorios") == 3)
    ok("10 unidades por mês (30 pendências)", contar(pg, "pendencias") == 30)
    ir(pg, url, "#/importar", 300); pg.set_input_files("#arqMes", str(ex / "powerbi1-2026-08.csv")); pg.wait_for_timeout(800)
    ok("agosto acusa 1 unidade fora do cadastro (RN03)", pg.locator("#qtdDesconhecidas").inner_text() == "1 fora do cadastro")
    pg.select_option("select[id^=assoc-]", "Regional Sinop"); pg.click("[data-assoc]"); pg.wait_for_timeout(500)
    ok("associar grava o apelido na unidade", "Regional de Sinop (nova sigla)" in pg.evaluate("async ()=> JSON.stringify((await db.get('unidades','Regional Sinop')).apelidos)"))
    ir(pg, url, "#/importar", 300); pg.set_input_files("#arqMes", str(ex / "powerbi1-2026-08.csv")); pg.wait_for_timeout(800)
    ok("na próxima importação, o apelido casa sozinho", pg.locator("#qtdDesconhecidas").count() == 0)
    ok("total do relatório guardado (setembro)", pg.evaluate("async ()=> (await db.get('relatorios','2026-09')).total_relatorio") is not None)
    # Excel e arquivo sem mês
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.append(["Mês", "Unidade", "Laudos pendentes"]); ws.append(["10/2026", "IML Cuiabá", 321]); ws.append(["10/2026", "Total", 321])
    xl = TMP / "relatorio.xlsx"; wb.save(xl)
    ir(pg, url, "#/importar", 300); pg.set_input_files("#arqMes", str(xl)); pg.wait_for_timeout(800)
    ok("Excel com mês 10/2026 é lido como 2026-10", pg.evaluate("async ()=> (await db.get('relatorios','2026-10'))?.tipo") == "excel")
    sem = TMP / "sem-mes.csv"; sem.write_text("Unidade;Laudos pendentes\nIML Cuiabá;10\n", encoding="utf-8")
    ir(pg, url, "#/importar", 300); pg.set_input_files("#arqMes", str(sem)); pg.wait_for_timeout(500)
    ok("arquivo sem mês é recusado (RN01)", "RN01" in pg.inner_text("#errosMes"))


# ------------------------------------------------------------------ CIN segunda via (André, Flávia e Glauce)
@cenario("cin-segunda-via")
def s2_pedido_do_cidadao(pg, url, ok):
    pg.set_viewport_size({"width": 390, "height": 800})   # celular
    ir(pg, url, "#/pedido", 300)
    pg.check("input[name=valida][value=nao]"); pg.check("input[name=mudou][value=nao]"); pg.click("#formElegivel button[type=submit]"); pg.wait_for_timeout(200)
    ok("CIN vencida leva ao agendamento no posto (RN03)", pg.locator("#avisoPosto a[href*='agendamento-cin']").count() == 1)
    pg.check("input[name=valida][value=sim]"); pg.click("#formElegivel button[type=submit]"); pg.wait_for_timeout(200)
    ok("elegível abre o formulário", pg.locator("#formPedido").count() == 1)
    pg.fill("#cNome", "Alice Teste"); pg.fill("#cCpf", "000.000.010-28"); pg.fill("#cNasc", "1990-05-10"); pg.fill("#cValidade", "2030-01-01"); pg.fill("#cContato", "alice@exemplo.com")
    pg.click("#formPedido button[type=submit]"); pg.wait_for_timeout(300)
    erros = pg.inner_text("#errosPedido")
    ok("CPF inválido é recusado (RN01)", "CPF inválido" in erros)
    ok("sem fotos o pedido não segue (RN02)", "foto do rosto" in erros and "foto do documento" in erros)
    valido = pg.evaluate("()=> dadosEmbutidos('seed').pedidos[0].cpf")
    pg.fill("#cCpf", valido)
    pg.click("[data-teste=rosto]"); pg.click("[data-teste=documento]"); pg.wait_for_timeout(800)
    pg.click("#formPedido button[type=submit]"); pg.wait_for_timeout(600)
    ok("pedido completo gera protocolo", pg.locator("#numProtocolo").count() == 1 and pg.inner_text("#numProtocolo").startswith("2V-"))
    ok("pedido gravado como recebido", pg.evaluate("async ()=> (await db.list('pedidos')).filter(p=>p.situacao==='recebido').length") == 1)
    ok("2 fotos gravadas, reduzidas para até 800 px", pg.evaluate("""async ()=>{ const fs = await db.list('fotos'); if (fs.length !== 2) return false;
        for (const f of fs){ const b = await createImageBitmap(f.conteudo); if (Math.max(b.width, b.height) > 800) return false; } return true; }"""))
    ok("sem rolagem horizontal no celular", pg.evaluate("()=> document.documentElement.scrollWidth <= window.innerWidth"))


# ------------------------------------------------------------------ TAP (Danielle)
@cenario("tap-portfolio-projetos")
def s2_tap_guiado(pg, url, ok):
    perfil(pg, url, "demandante")
    pg.evaluate("async ()=>{ for (const d of dadosEmbutidos('seed').documentos) await db.put('documentos', d); }")
    ir(pg, url, "#/tap/novo", 500)
    ok("TAP vazio mostra as pendências", int(pg.inner_text("#qtdPend")) >= 9)
    pg.fill("#tap-titulo", "Projeto de teste"); pg.fill("#tap-unidade", "Laboratório Forense"); pg.fill("#tap-responsavel", "Responsável Teste")
    pg.fill("#tap-contato", "teste@politec.mt.gov.br"); pg.wait_for_timeout(400)
    pg.reload(); pg.wait_for_timeout(600)
    ok("rascunho salvo: o título volta depois de recarregar", pg.input_value("#tap-titulo") == "Projeto de teste")
    pg.click("[data-etapa=objetivo]"); pg.wait_for_timeout(300)
    pg.fill("#tap-objetivo", "o" * 120); pg.fill("#tap-justificativa", "j" * 150); pg.wait_for_timeout(300)
    ok("justificativa curta segue pendente (RN02)", "Justificativa: 150 de 200" in pg.inner_text("#painelPendencias"))
    pg.fill("#tap-justificativa", "j" * 210); pg.wait_for_timeout(300)
    pg.click("[data-etapa=alinhamento]"); pg.wait_for_timeout(300); pg.select_option("#tap-objetivo_estrategico", index=1); pg.wait_for_timeout(200)
    pg.click("[data-etapa=custo]"); pg.wait_for_timeout(300); pg.fill("#tap-valor_estimado", "150.000,00"); pg.fill("#tap-prazo_meses", "12"); pg.wait_for_timeout(300)
    pg.click("[data-etapa=anexos]"); pg.wait_for_timeout(300)
    ok("faltam os 3 anexos obrigatórios", pg.inner_text("#qtdPend") == "3")
    pg.set_input_files("[data-anexo=estudo_tecnico]", {"name": "grande.pdf", "mimeType": "application/pdf", "buffer": b"0" * (11 * 1024 * 1024)}); pg.wait_for_timeout(400)
    ok("anexo de 11 MB é recusado (RN04)", contar(pg, "anexos") == 0)
    for t in ("pesquisa_precos", "estudo_tecnico"):
        pg.set_input_files(f"[data-anexo={t}]", {"name": f"{t}.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4"}); pg.wait_for_timeout(500)
    ok("2 anexos guardados", contar(pg, "anexos") == 2)
    pg.fill("#aus-anuencia_chefia", "A chefia assina no SEI depois da triagem."); pg.wait_for_timeout(400)
    ok("ausência justificada fecha a última pendência (RN03)", pg.inner_text("#qtdPend") == "0")
    pg.click("#btnConcluir"); pg.wait_for_timeout(500)
    ok("TAP concluído fica pronto para envio", pg.evaluate("async ()=> (await db.list('propostas')).filter(p=>p.status==='pronto para envio').length") == 1)
    ok("valor estimado gravado em reais sem erro", pg.evaluate("async ()=> (await db.list('propostas')).find(p=>p.status==='pronto para envio').valor_estimado") == 150000)



@cenario("mapa-distribuicao-credor")
def s3_mapa_por_item(pg, url, ok):
    from decimal import Decimal
    from openpyxl import load_workbook
    ex = APPS / "mapa-distribuicao-credor" / "exemplos"
    ir(pg, url, "#/processos"); pg.click("#btnExemplo"); pg.wait_for_timeout(800)
    ir(pg, url, "#/mapa/PROC-2026-001", 600)
    gab = load_workbook(ex / "mapa-manual-referencia.xlsx", data_only=True)["2 Por órgão"]
    por_orgao = {}
    for orgao, _g, _i, _q, est, _h in gab.iter_rows(min_row=2, values_only=True):
        por_orgao[orgao] = por_orgao.get(orgao, Decimal("0")) + Decimal(str(est))
    app = pg.evaluate("async ()=>{ const s = dadosEmbutidos('seed'); const m = montarMapa(await db.porIndice('itens','id_processo','PROC-2026-001'), await db.porIndice('distribuicao','id_processo','PROC-2026-001')); return Object.fromEntries(Object.entries(m.geral.porOrgao).map(([o,x])=>[o,x.v])); }")
    ok("estimado por órgão igual ao gabarito, centavo a centavo", all(int((por_orgao[o] * 100).to_integral_value()) == app[o] for o in por_orgao))
    total = sum(por_orgao.values())
    ok("total geral na tela igual ao gabarito", pg.inner_text("#totalGeral").replace(" ", " ") == "R$ " + f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    ok("6 subtotais de grupo e 1 total geral", pg.locator("tr.subtotal").count() == 6 and pg.locator("tr.total").count() == 1)
    pg.locator("#tabelaMapa [data-origem]").first.click(); pg.wait_for_timeout(200)
    ok("clicar num valor mostra arquivo, aba e linha", "distribuicao-orgaos.xlsx" in pg.inner_text("#origemValor") and "linha" in pg.inner_text("#origemValor"))


@cenario("equipamentos-google-earth")
def s3_equipamentos_e_fotos(pg, url, ok):
    import io
    import random
    from PIL import Image
    perfil(pg, url, "coordenadoria")
    ir(pg, url, "#/config"); pg.click("#btnSeed"); pg.wait_for_timeout(800)
    ir(pg, url, "#/equipamentos", 500)
    ok("lista mostra os 150 equipamentos", pg.inner_text("#qtdEq") == "150")
    pg.select_option("#fCrit", "5"); pg.wait_for_timeout(300)
    n5 = pg.evaluate("async ()=> (await db.list('equipamentos')).filter(e=>e.criticidade===5).length")
    ok("filtro de criticidade soma o total certo", pg.inner_text("#qtdEq") == str(n5))
    ir(pg, url, "#/equipamento/novo", 500)
    uni = pg.evaluate("async ()=> (await db.list('unidades'))[0].id")
    pat = pg.evaluate("async ()=> (await db.list('equipamentos'))[0].patrimonio")
    pg.fill("#ePat", pat); pg.fill("#eNome", "Cromatógrafo teste"); pg.fill("#eTipo", "Cromatógrafo"); pg.select_option("#eUnidade", uni)
    pg.click("#formEq button[type=submit]"); pg.wait_for_timeout(400)
    ok("patrimônio duplicado é apontado", "já está cadastrado" in pg.inner_text("#errosEq"))
    pg.fill("#ePat", "PAT-TESTE-1"); pg.click("#formEq button[type=submit]"); pg.wait_for_timeout(600)
    ok("equipamento salvo abre a ficha", "#/equipamento/" in pg.url and "PAT-TESTE-1" in pg.inner_text("h1"))
    fotos = []
    for k in range(3):   # fotos de celular de cerca de 4 MB (ruído quase não comprime)
        im = Image.frombytes("RGB", (4000, 3000), random.randbytes(4000 * 3000 * 3)); b = io.BytesIO(); im.save(b, "JPEG", quality=85)
        f = TMP / f"foto{k}.jpg"; f.write_bytes(b.getvalue()); fotos.append(str(f))
    ok("fotos de teste com mais de 3 MB cada", all(pathlib.Path(f).stat().st_size > 3_000_000 for f in fotos))
    pg.set_input_files("#inpFoto", fotos); pg.wait_for_timeout(8000)
    ok("a ficha mostra as 3 fotos", pg.locator("#galeria img").count() == 3)
    tam = pg.evaluate("async ()=> (await db.list('arquivos')).reduce((a,x)=>a+x.tamanho,0)")
    ok(f"as 3 fotos ocupam menos de 1 MB ({tam // 1024} KB)", tam < 1024 * 1024)
    ir(pg, url, "#/equipamentos", 400)
    with pg.expect_download() as dl:
        pg.click("#btnCsvEq")
    csv = pathlib.Path(dl.value.path()).read_bytes()
    ok("CSV com BOM e ponto e vírgula", csv.startswith(b"\xef\xbb\xbf") and b"patrimonio;nome" in csv)



@cenario("comissao-etica-denuncias")
def s2_registro_da_denuncia(pg, url, ok):
    perfil(pg, url, "secretario")
    ir(pg, url, "#/denuncia/nova", 400)
    pg.fill("#dDenunciante", "Denunciante Teste"); pg.fill("#dDenunciado", "Servidor Teste"); pg.fill("#dRelato", "curto")
    pg.click("#formDenuncia button[type=submit]"); pg.wait_for_timeout(300)
    ok("relato curto é recusado", "50 caracteres" in pg.inner_text("#errosDenuncia") and contar(pg, "denuncias") == 0)
    pg.fill("#dRelato", "O chefe gritou com a servidora na frente da equipe e a humilhou, com ameaças de remoção do setor.")
    pg.set_input_files("#dAnexos", {"name": "print.png", "mimeType": "image/png", "buffer": b"PNG teste"})
    pg.click("#formDenuncia button[type=submit]"); pg.wait_for_timeout(600)
    ok("denúncia registrada com protocolo CE-AAAA-001", pg.evaluate("async ()=> (await db.list('denuncias'))[0]?.protocolo || ''").endswith("-001"))
    ok("anexo guardado", contar(pg, "anexos") == 1)
    ok("histórico registra a ação com o perfil (RN05)", "denúncia registrada" in pg.inner_text("#historico") and "secretario" in pg.inner_text("#historico"))
    ir(pg, url, "#/painel", 300)
    ok("o caso aparece no painel, na etapa registro", pg.inner_text("[data-etapa-qtd=registro]") == "1")


# cenários dos sprints seguintes ficam em _base/cenarios/sprintN.py (mesmas funções e o mesmo @cenario)
for _arq in sorted((pathlib.Path(__file__).parent / "cenarios").glob("*.py")):
    exec(compile(_arq.read_text(encoding="utf-8"), str(_arq), "exec"), globals())


def main():
    alvos = sys.argv[1:] or list(CENARIOS)
    erro = False
    with sync_playwright() as p:
        nav = p.chromium.launch()
        for app in alvos:
            for fn in CENARIOS.get(app, []):
                ctx = nav.new_context(accept_downloads=True); pg = ctx.new_page()
                console = []
                pg.on("pageerror", lambda e: console.append(str(e)))
                pg.on("console", lambda m: console.append(m.text) if m.type == "error" else None)
                pg.on("dialog", lambda d: d.accept())
                resultados = []
                def ok(nome, cond):
                    resultados.append((nome, bool(cond)))
                try:
                    fn(pg, (APPS / app / "index.html").as_uri(), ok)
                except Exception as e:
                    resultados.append((f"exceção: {str(e).splitlines()[0][:160]}", False))
                if console:
                    resultados.append((f"erros no console: {console[:2]}", False))
                falhas = [n for n, c in resultados if not c]
                print(("OK   " if not falhas else "FALHA"), app, "·", fn.__name__, f"· {len(resultados) - len(falhas)}/{len(resultados)}")
                for n in falhas:
                    print("      -", n)
                erro |= bool(falhas)
                ctx.close()
        nav.close()
    sys.exit(1 if erro else 0)


if __name__ == "__main__":
    main()
